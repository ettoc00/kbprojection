import argparse
import json
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from kbprojection.settings import get_default_results_dir


PROMPT_STYLES = ("legacy_icl", "icl", "lasha")
LEGACY_STATUS_MAP = {
    "already_correct": "baseline_solved",
    "fixed": "normalised_kb_solved",
    "fixed_raw_kb": "raw_kb_solved",
    "still_wrong": "kb_not_solved",
    "still_wrong_raw_kb": "kb_not_solved",
    "empty_kb_after_filter": "kb_normalisation_empty",
    "llm_error": "kb_generation_failed",
}


def _json_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    return json.dumps(value, ensure_ascii=False, indent=2, default=str)


def _premises_text(problem: dict[str, Any]) -> str:
    return "\n".join(str(p) for p in (problem.get("premises") or []))


def _extract_prover_call(item: dict[str, Any], field: str, fallback_index: int) -> dict[str, Any]:
    call = item.get(field)
    if isinstance(call, dict):
        return call
    calls = item.get("prover_calls") or []
    if fallback_index < len(calls) and isinstance(calls[fallback_index], dict):
        return calls[fallback_index]
    return {}


def _proof_text(call: dict[str, Any]) -> str:
    rendered = call.get("proofs_rendered")
    if isinstance(rendered, dict) and rendered:
        parts = []
        for label in ("entailment", "contradiction"):
            text = rendered.get(label)
            if text:
                parts.append(f"[{label}]\n{text}")
        for label, text in rendered.items():
            if label not in {"entailment", "contradiction"} and text:
                parts.append(f"[{label}]\n{text}")
        if parts:
            return "\n\n".join(parts)
    return _json_text(call.get("proofs"))


def _normalized_final_status(value: str) -> str:
    return LEGACY_STATUS_MAP.get(value, value)


def _compute_current_status(
    item: dict[str, Any],
    gold_label: str,
    baseline_call: dict[str, Any],
    raw_call: dict[str, Any],
    norm_call: dict[str, Any],
) -> str:
    if baseline_call.get("error"):
        return "baseline_prover_failed"

    if baseline_call.get("label") == gold_label:
        return "baseline_solved"

    if item.get("llm_error"):
        return "kb_generation_failed"

    kb_raw = item.get("kb_raw") or []
    if not kb_raw:
        return "kb_generation_empty"

    if raw_call.get("label") == gold_label:
        return "raw_kb_solved"

    kb_filtered = item.get("kb_filtered") or []
    if not kb_filtered:
        return "kb_normalisation_empty"

    if norm_call.get("error"):
        return "normalised_kb_prover_failed"

    if norm_call.get("label") == gold_label:
        return "normalised_kb_solved"

    return "kb_not_solved"


def _build_row(model: str, prompt_style: str, item: dict[str, Any]) -> dict[str, Any]:
    problem = item.get("problem") or {}
    baseline_call = _extract_prover_call(item, "prover_call_no_kb", 0)
    raw_call = _extract_prover_call(item, "prover_call_raw_kb", 1)
    norm_call = _extract_prover_call(item, "prover_call_kb", 2)
    gold_label = problem.get("gold_label", "")
    original_final_status = _normalized_final_status(item.get("final_status", ""))
    current_final_status = _compute_current_status(item, gold_label, baseline_call, raw_call, norm_call)
    found_solution = current_final_status in {"baseline_solved", "raw_kb_solved", "normalised_kb_solved"}
    baseline_solved_current = baseline_call.get("label") == gold_label if baseline_call else False
    raw_solved_current = raw_call.get("label") == gold_label if raw_call else False
    normalised_solved_current = norm_call.get("label") == gold_label if norm_call else False
    baseline_solved_saved = original_final_status == "baseline_solved"
    problem_solved = found_solution or baseline_solved_current or baseline_solved_saved
    kb_success_type = (
        "baseline"
        if current_final_status == "baseline_solved"
        else "raw_kb"
        if current_final_status == "raw_kb_solved"
        else "normalised_kb"
        if current_final_status == "normalised_kb_solved"
        else ""
    )

    return {
        "model": model,
        "prompt_style": prompt_style,
        "problem_id": problem.get("id", ""),
        "dataset": problem.get("dataset", ""),
        "split": problem.get("split", ""),
        "gold_label": problem.get("gold_label", ""),
        "premises": _premises_text(problem),
        "hypothesis": problem.get("hypothesis", ""),
        "pred_no_kb": item.get("pred_no_kb", ""),
        "status_no_kb": item.get("status_no_kb", ""),
        "pred_with_raw_kb": item.get("pred_with_raw_kb", ""),
        "status_with_raw_kb": item.get("status_with_raw_kb", ""),
        "pred_with_kb": item.get("pred_with_kb", ""),
        "status_with_kb": item.get("status_with_kb", ""),
        "final_status": original_final_status,
        "current_final_status": current_final_status,
        "found_solution": found_solution,
        "found_solution_int": 1 if found_solution else 0,
        "problem_solved": problem_solved,
        "problem_solved_int": 1 if problem_solved else 0,
        "fixed_by": item.get("fixed_by", ""),
        "baseline_solved_current": baseline_solved_current,
        "baseline_solved_int": 1 if baseline_solved_current else 0,
        "baseline_solved_saved": baseline_solved_saved,
        "baseline_solved_saved_int": 1 if baseline_solved_saved else 0,
        "raw_solved_current": raw_solved_current,
        "raw_solved_int": 1 if raw_solved_current else 0,
        "normalised_solved_current": normalised_solved_current,
        "normalised_solved_int": 1 if normalised_solved_current else 0,
        "kb_success_type": kb_success_type,
        "kb_raw": _json_text(item.get("kb_raw")),
        "kb_filtered": _json_text(item.get("kb_filtered")),
        "kb_details": _json_text(item.get("kb_details")),
        "llm_error": item.get("llm_error", "") or "",
        "llm_output_raw": item.get("llm_output_raw", "") or "",
        "baseline_prover_label": baseline_call.get("label", ""),
        "baseline_prover_error": baseline_call.get("error", ""),
        "baseline_prover_tree": _proof_text(baseline_call),
        "raw_prover_label": raw_call.get("label", ""),
        "raw_prover_error": raw_call.get("error", ""),
        "raw_prover_tree": _proof_text(raw_call),
        "normalised_prover_label": norm_call.get("label", ""),
        "normalised_prover_error": norm_call.get("error", ""),
        "normalised_prover_tree": _proof_text(norm_call),
        "prover_backfill_mismatches": _json_text(item.get("prover_backfill_mismatches")),
    }


def _load_compare_rows(path: Path) -> list[dict[str, Any]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    model = payload.get("meta", {}).get("model") or ""
    rows = []
    for prompt_style, run in (payload.get("runs") or {}).items():
        for item in (run.get("items_by_id") or {}).values():
            tagged = dict(item)
            tagged["_source_file"] = str(path)
            rows.append(_build_row(model, prompt_style, tagged))
    return rows


def _load_replay_rows(path: Path) -> list[dict[str, Any]]:
    items = json.loads(path.read_text(encoding="utf-8"))
    rows = []
    for item in items:
        tagged = dict(item)
        tagged["_source_file"] = str(path)
        rows.append(_build_row(tagged.get("model", ""), tagged.get("prompt_style", ""), tagged))
    return rows


def load_all_rows(results_dir: Path) -> list[dict[str, Any]]:
    compare_rows: list[dict[str, Any]] = []
    replay_rows: list[dict[str, Any]] = []
    entailment_scope_rows: list[dict[str, Any]] = []
    for path in sorted(results_dir.glob("2x2_50_compare__openrouter__*.json")):
        if "__backup_" not in path.name:
            compare_rows.extend(_load_compare_rows(path))
    for path in sorted(results_dir.glob("entailment_scope_compare__openrouter__*.json")):
        if "__backup_" not in path.name:
            entailment_scope_rows.extend(_load_compare_rows(path))
    for path in sorted(results_dir.glob("results_replay_openrouter_*__*.json")):
        if "__backup_" not in path.name:
            replay_rows.extend(_load_replay_rows(path))
    compare_models = {row["model"] for row in compare_rows}
    replay_models = {row["model"] for row in replay_rows}
    keep_models = compare_models & replay_models
    return entailment_scope_rows + [row for row in compare_rows + replay_rows if row["model"] in keep_models]


def _write_data_sheet(workbook: Workbook, rows: list[dict[str, Any]]) -> None:
    ws = workbook.active
    ws.title = "data"
    columns = [
        "model",
        "prompt_style",
        "problem_id",
        "dataset",
        "split",
        "gold_label",
        "premises",
        "hypothesis",
        "pred_no_kb",
        "status_no_kb",
        "pred_with_raw_kb",
        "status_with_raw_kb",
        "pred_with_kb",
        "status_with_kb",
        "final_status",
        "current_final_status",
        "found_solution",
        "found_solution_int",
        "problem_solved",
        "problem_solved_int",
        "fixed_by",
        "baseline_solved_current",
        "baseline_solved_int",
        "baseline_solved_saved",
        "baseline_solved_saved_int",
        "raw_solved_current",
        "raw_solved_int",
        "normalised_solved_current",
        "normalised_solved_int",
        "kb_success_type",
        "kb_raw",
        "kb_filtered",
        "kb_details",
        "llm_error",
        "llm_output_raw",
        "baseline_prover_label",
        "baseline_prover_error",
        "baseline_prover_tree",
        "raw_prover_label",
        "raw_prover_error",
        "raw_prover_tree",
        "normalised_prover_label",
        "normalised_prover_error",
        "normalised_prover_tree",
        "prover_backfill_mismatches",
    ]
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    column_index = {name: idx + 1 for idx, name in enumerate(columns)}
    current_status_col = get_column_letter(column_index["current_final_status"])
    final_status_col = get_column_letter(column_index["final_status"])
    found_solution_col = get_column_letter(column_index["found_solution"])
    found_solution_int_col = get_column_letter(column_index["found_solution_int"])
    problem_solved_col = get_column_letter(column_index["problem_solved"])
    problem_solved_int_col = get_column_letter(column_index["problem_solved_int"])

    status_logic_status_col = "$A:$A"
    status_logic_found_col = "$B:$B"
    status_logic_problem_col = "$C:$C"

    for row in sorted(rows, key=lambda r: (r["model"], r["prompt_style"], r["problem_id"])):
        values = [row.get(column, "") for column in columns]
        ws.append(values)
        row_idx = ws.max_row
        ws[f"{found_solution_col}{row_idx}"] = (
            f'=IFERROR(INDEX(status_logic!{status_logic_found_col},'
            f'MATCH({current_status_col}{row_idx},status_logic!{status_logic_status_col},0)),FALSE)'
        )
        ws[f"{found_solution_int_col}{row_idx}"] = f'=--{found_solution_col}{row_idx}'
        ws[f"{problem_solved_col}{row_idx}"] = (
            f'=OR({final_status_col}{row_idx}="baseline_solved",'
            f'IFERROR(INDEX(status_logic!{status_logic_problem_col},'
            f'MATCH({current_status_col}{row_idx},status_logic!{status_logic_status_col},0)),FALSE))'
        )
        ws[f"{problem_solved_int_col}{row_idx}"] = f'=--{problem_solved_col}{row_idx}'
    for idx, column in enumerate(columns, start=1):
        max_len = len(column)
        for row_idx in range(2, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=idx).value
            if value is None:
                continue
            max_len = min(80, max(max_len, len(str(value))))
        ws.column_dimensions[get_column_letter(idx)].width = max_len + 2
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _write_status_logic_sheet(workbook: Workbook) -> None:
    ws = workbook.create_sheet("status_logic")
    columns = ["status", "found_solution", "problem_solved"]
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    rows = [
        ("baseline_solved", False, True),
        ("raw_kb_solved", True, True),
        ("normalised_kb_solved", True, True),
        ("kb_not_solved", False, False),
        ("kb_generation_failed", False, False),
        ("kb_generation_empty", False, False),
        ("kb_normalisation_empty", False, False),
        ("baseline_prover_failed", False, False),
        ("normalised_kb_prover_failed", False, False),
    ]
    for row in rows:
        ws.append(list(row))
    for idx, column in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = max(len(column), 18)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _write_problem_view_sheet(workbook: Workbook, rows: list[dict[str, Any]]) -> None:
    ws = workbook.create_sheet("problem_view")
    columns = [
        "data_row",
        "model",
        "prompt_style",
        "problem_id",
        "dataset",
        "split",
        "gold_label",
        "current_final_status",
        "found_solution",
        "problem_solved",
        "legacy_icl_problem_solved_for_problem",
        "icl_problem_solved_for_problem",
        "lasha_problem_solved_for_problem",
        "best_prompt_for_problem",
        "best_problem_solved_count",
        "premises",
        "hypothesis",
    ]
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    total_rows = len(rows)
    for data_row in range(2, total_rows + 2):
        ws.append([""] * len(columns))
        row_idx = ws.max_row
        ws[f"A{row_idx}"] = data_row
        ws[f"B{row_idx}"] = f"=data!A{data_row}"
        ws[f"C{row_idx}"] = f"=data!B{data_row}"
        ws[f"D{row_idx}"] = f"=data!C{data_row}"
        ws[f"E{row_idx}"] = f'=INDEX(data!$D:$D,MATCH($D{row_idx},data!$C:$C,0))'
        ws[f"F{row_idx}"] = f'=INDEX(data!$E:$E,MATCH($D{row_idx},data!$C:$C,0))'
        ws[f"G{row_idx}"] = f'=INDEX(data!$F:$F,MATCH($D{row_idx},data!$C:$C,0))'
        ws[f"H{row_idx}"] = f"=data!P{data_row}"
        ws[f"I{row_idx}"] = f"=data!Q{data_row}"
        ws[f"J{row_idx}"] = f"=data!S{data_row}"
        ws[f"K{row_idx}"] = f'=COUNTIFS(data!$C:$C,$D{row_idx},data!$B:$B,"legacy_icl",data!$T:$T,1)'
        ws[f"L{row_idx}"] = f'=COUNTIFS(data!$C:$C,$D{row_idx},data!$B:$B,"icl",data!$T:$T,1)'
        ws[f"M{row_idx}"] = f'=COUNTIFS(data!$C:$C,$D{row_idx},data!$B:$B,"lasha",data!$T:$T,1)'
        ws[f"O{row_idx}"] = f"=MAX(K{row_idx}:M{row_idx})"
        ws[f"N{row_idx}"] = (
            f'=IF(K{row_idx}=O{row_idx},"legacy_icl",'
            f'IF(L{row_idx}=O{row_idx},"icl","lasha"))'
        )
        ws[f"P{row_idx}"] = f'=INDEX(data!$G:$G,MATCH($D{row_idx},data!$C:$C,0))'
        ws[f"Q{row_idx}"] = f'=INDEX(data!$H:$H,MATCH($D{row_idx},data!$C:$C,0))'

    for idx, column in enumerate(columns, start=1):
        max_len = len(column)
        for row_idx in range(2, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=idx).value
            if value is None:
                continue
            max_len = min(80, max(max_len, len(str(value))))
        ws.column_dimensions[get_column_letter(idx)].width = max_len + 2
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _write_summary_sheet(workbook: Workbook) -> None:
    ws = workbook.create_sheet("summary")
    columns = ["prompt_style", "found_solution", "problem_solved", "best_on_problem_view"]
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    prompt_map = [("legacy_icl", 2, "E", "H"), ("icl", 3, "F", "I"), ("lasha", 4, "G", "J")]
    for prompt_style, row_idx, found_col, solved_col in prompt_map:
        ws[f"A{row_idx}"] = prompt_style
        ws[f"B{row_idx}"] = f"=SUM(problem_view!${found_col}:${found_col})"
        ws[f"C{row_idx}"] = f"=SUM(problem_view!${solved_col}:${solved_col})"
        ws[f"D{row_idx}"] = f'=COUNTIF(problem_view!$N:$N,"{prompt_style}")'
    for idx, column in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = max(len(column), 20)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _is_raw_solved(row: dict[str, Any]) -> bool:
    return bool(row["raw_solved_current"])


def _is_norm_solved(row: dict[str, Any]) -> bool:
    return bool(row["normalised_solved_current"])


def _is_baseline_solved(row: dict[str, Any]) -> bool:
    return bool(row["baseline_solved_current"])


def _build_overview_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, Counter[str]] = defaultdict(Counter)
    for row in rows:
        prompt_style = row["prompt_style"]
        grouped[prompt_style]["rows"] += 1
        grouped[prompt_style]["baseline_solved"] += int(_is_baseline_solved(row))
        grouped[prompt_style]["raw_solved"] += int(_is_raw_solved(row))
        grouped[prompt_style]["norm_solved"] += int(_is_norm_solved(row))
        grouped[prompt_style][f"status::{row['current_final_status']}"] += 1
        grouped[prompt_style]["mismatch_rows"] += int(bool(row["prover_backfill_mismatches"]))

    overview = []
    for prompt_style in sorted(grouped):
        counts = grouped[prompt_style]
        total = counts["rows"]
        overview.append(
            {
                "prompt_style": prompt_style,
                "rows": total,
                "baseline_solved": counts["baseline_solved"],
                "raw_solved": counts["raw_solved"],
                "norm_solved": counts["norm_solved"],
                "baseline_rate": counts["baseline_solved"] / total if total else 0.0,
                "raw_rate": counts["raw_solved"] / total if total else 0.0,
                "norm_rate": counts["norm_solved"] / total if total else 0.0,
                "kb_not_solved": counts["status::kb_not_solved"],
                "kb_generation_failed": counts["status::kb_generation_failed"],
                "kb_generation_empty": counts["status::kb_generation_empty"],
                "kb_normalisation_empty": counts["status::kb_normalisation_empty"],
                "baseline_prover_failed": counts["status::baseline_prover_failed"],
                "normalised_kb_prover_failed": counts["status::normalised_kb_prover_failed"],
                "mismatch_rows": counts["mismatch_rows"],
            }
        )
    return overview


def _build_model_prompt_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[tuple[str, str], Counter[str]] = defaultdict(Counter)
    for row in rows:
        key = (row["model"], row["prompt_style"])
        grouped[key]["rows"] += 1
        grouped[key]["baseline_solved"] += int(_is_baseline_solved(row))
        grouped[key]["raw_solved"] += int(_is_raw_solved(row))
        grouped[key]["norm_solved"] += int(_is_norm_solved(row))
        grouped[key][f"status::{row['current_final_status']}"] += 1
        grouped[key]["mismatch_rows"] += int(bool(row["prover_backfill_mismatches"]))

    out = []
    for (model, prompt_style), counts in sorted(grouped.items()):
        total = counts["rows"]
        out.append(
            {
                "model": model,
                "prompt_style": prompt_style,
                "rows": total,
                "baseline_solved": counts["baseline_solved"],
                "raw_solved": counts["raw_solved"],
                "norm_solved": counts["norm_solved"],
                "baseline_rate": counts["baseline_solved"] / total if total else 0.0,
                "raw_rate": counts["raw_solved"] / total if total else 0.0,
                "norm_rate": counts["norm_solved"] / total if total else 0.0,
                "kb_not_solved": counts["status::kb_not_solved"],
                "kb_generation_failed": counts["status::kb_generation_failed"],
                "kb_generation_empty": counts["status::kb_generation_empty"],
                "kb_normalisation_empty": counts["status::kb_normalisation_empty"],
                "baseline_prover_failed": counts["status::baseline_prover_failed"],
                "normalised_kb_prover_failed": counts["status::normalised_kb_prover_failed"],
                "mismatch_rows": counts["mismatch_rows"],
            }
        )
    return out


def _build_problem_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, dict[str, Any]] = {}
    for row in rows:
        pid = row["problem_id"]
        bucket = grouped.setdefault(
            pid,
            {
                "problem_id": pid,
                "dataset": row["dataset"],
                "split": row["split"],
                "gold_label": row["gold_label"],
                "premises": row["premises"],
                "hypothesis": row["hypothesis"],
                "rows": 0,
                "baseline_solved": 0,
                "raw_solved": 0,
                "norm_solved": 0,
                "models": set(),
                "prompt_styles": set(),
                "statuses": Counter(),
                "norm_by_prompt": Counter(),
                "raw_by_prompt": Counter(),
            },
        )
        bucket["rows"] += 1
        bucket["baseline_solved"] += int(_is_baseline_solved(row))
        bucket["raw_solved"] += int(_is_raw_solved(row))
        bucket["norm_solved"] += int(_is_norm_solved(row))
        bucket["models"].add(row["model"])
        bucket["prompt_styles"].add(row["prompt_style"])
        bucket["statuses"][row["current_final_status"]] += 1
        if _is_raw_solved(row):
            bucket["raw_by_prompt"][row["prompt_style"]] += 1
        if _is_norm_solved(row):
            bucket["norm_by_prompt"][row["prompt_style"]] += 1

    out = []
    for pid, bucket in sorted(grouped.items()):
        row = {
            "problem_id": pid,
            "dataset": bucket["dataset"],
            "split": bucket["split"],
            "gold_label": bucket["gold_label"],
            "rows": bucket["rows"],
            "models": len(bucket["models"]),
            "prompt_styles": len(bucket["prompt_styles"]),
            "baseline_solved": bucket["baseline_solved"],
            "raw_solved": bucket["raw_solved"],
            "norm_solved": bucket["norm_solved"],
            "kb_not_solved": bucket["statuses"]["kb_not_solved"],
            "kb_generation_failed": bucket["statuses"]["kb_generation_failed"],
            "kb_generation_empty": bucket["statuses"]["kb_generation_empty"],
            "kb_normalisation_empty": bucket["statuses"]["kb_normalisation_empty"],
            "baseline_prover_failed": bucket["statuses"]["baseline_prover_failed"],
            "normalised_kb_prover_failed": bucket["statuses"]["normalised_kb_prover_failed"],
            "premises": bucket["premises"],
            "hypothesis": bucket["hypothesis"],
        }
        for prompt_style in PROMPT_STYLES:
            row[f"{prompt_style}_raw_solved"] = bucket["raw_by_prompt"][prompt_style]
            row[f"{prompt_style}_norm_solved"] = bucket["norm_by_prompt"][prompt_style]
        out.append(row)
    return out


def _parse_kb_list(text: str) -> list[str]:
    if not text:
        return []
    try:
        value = json.loads(text)
    except Exception:
        return []
    if isinstance(value, list):
        return [str(v) for v in value]
    return []


def _build_kb_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[tuple[str, str], dict[str, Any]] = {}
    for row in rows:
        if not row["found_solution"]:
            continue
        success_stage = row["kb_success_type"]
        if success_stage == "baseline":
            continue
        kb_items = _parse_kb_list(row["kb_raw"] if success_stage == "raw_kb" else row["kb_filtered"])
        for relation in kb_items:
            key = (relation, success_stage)
            bucket = grouped.setdefault(
                key,
                {
                    "kb_relation": relation,
                    "success_stage": success_stage,
                    "successful_rows": 0,
                    "models": set(),
                    "prompt_styles": set(),
                    "problems": set(),
                },
            )
            bucket["successful_rows"] += 1
            bucket["models"].add(row["model"])
            bucket["prompt_styles"].add(row["prompt_style"])
            bucket["problems"].add(row["problem_id"])

    out = []
    for (_relation, _stage), bucket in sorted(
        grouped.items(),
        key=lambda item: (-item[1]["successful_rows"], item[0][0], item[0][1]),
    ):
        out.append(
            {
                "kb_relation": bucket["kb_relation"],
                "success_stage": bucket["success_stage"],
                "successful_rows": bucket["successful_rows"],
                "distinct_models": len(bucket["models"]),
                "distinct_prompt_styles": len(bucket["prompt_styles"]),
                "distinct_problems": len(bucket["problems"]),
            }
        )
    return out


def _build_prompt_success_matrix(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for prompt_style in PROMPT_STYLES:
        sub = [row for row in rows if row["prompt_style"] == prompt_style]
        out.append(
            {
                "prompt_style": prompt_style,
                "problem_solved": sum(row["problem_solved_int"] for row in sub),
                "kb_found_solution": sum(row["found_solution_int"] for row in sub),
                "raw_solved": sum(row["raw_solved_int"] for row in sub),
                "normalised_solved": sum(row["normalised_solved_int"] for row in sub),
                "baseline_solved": sum(row["baseline_solved_int"] for row in sub),
            }
        )
    return out


def _build_model_prompt_matrix(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, dict[str, int]] = defaultdict(dict)
    for model in sorted({row["model"] for row in rows}):
        grouped[model] = {"model": model}
        for prompt_style in PROMPT_STYLES:
            sub = [row for row in rows if row["model"] == model and row["prompt_style"] == prompt_style]
            grouped[model][f"{prompt_style}_problem_solved"] = sum(row["problem_solved_int"] for row in sub)
            grouped[model][f"{prompt_style}_kb_found_solution"] = sum(row["found_solution_int"] for row in sub)
            grouped[model][f"{prompt_style}_raw_solved"] = sum(row["raw_solved_int"] for row in sub)
            grouped[model][f"{prompt_style}_normalised_solved"] = sum(row["normalised_solved_int"] for row in sub)
        scores = {prompt: grouped[model][f"{prompt}_problem_solved"] for prompt in PROMPT_STYLES}
        best_prompt = max(scores, key=scores.get)
        grouped[model]["best_prompt"] = best_prompt
        grouped[model]["best_prompt_score"] = scores[best_prompt]
    return [grouped[model] for model in sorted(grouped)]


def _build_problem_prompt_matrix(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, dict[str, Any]] = {}
    for problem_id in sorted({row["problem_id"] for row in rows}):
        sample = next(row for row in rows if row["problem_id"] == problem_id)
        grouped[problem_id] = {
            "problem_id": problem_id,
            "dataset": sample["dataset"],
            "split": sample["split"],
            "gold_label": sample["gold_label"],
            "premises": sample["premises"],
            "hypothesis": sample["hypothesis"],
        }
        for prompt_style in PROMPT_STYLES:
            sub = [row for row in rows if row["problem_id"] == problem_id and row["prompt_style"] == prompt_style]
            grouped[problem_id][f"{prompt_style}_problem_solved"] = sum(row["problem_solved_int"] for row in sub)
            grouped[problem_id][f"{prompt_style}_kb_found_solution"] = sum(row["found_solution_int"] for row in sub)
        scores = {prompt: grouped[problem_id][f"{prompt}_problem_solved"] for prompt in PROMPT_STYLES}
        best_prompt = max(scores, key=scores.get)
        grouped[problem_id]["best_prompt"] = best_prompt
        grouped[problem_id]["best_prompt_score"] = scores[best_prompt]
    return [grouped[problem_id] for problem_id in sorted(grouped)]


def _build_failure_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    failures = [
        row for row in rows
        if row["current_final_status"] not in {"baseline_solved", "raw_kb_solved", "normalised_kb_solved"}
    ]
    return sorted(failures, key=lambda r: (r["current_final_status"], r["model"], r["prompt_style"], r["problem_id"]))


def _build_mismatch_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    mismatches = [row for row in rows if row["prover_backfill_mismatches"]]
    return sorted(mismatches, key=lambda r: (r["model"], r["prompt_style"], r["problem_id"]))


def _build_base_workbook(rows: list[dict[str, Any]], output_path: Path) -> None:
    workbook = Workbook()
    _write_data_sheet(workbook, rows)
    _write_status_logic_sheet(workbook)
    _write_problem_view_sheet(workbook, rows)
    _write_summary_sheet(workbook)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export benchmark results to one data sheet and pivot sheets.")
    parser.add_argument(
        "--results-dir",
        default=str(get_default_results_dir()),
        help="Directory containing result JSON files.",
    )
    parser.add_argument(
        "--output",
        default=str(Path("scripts/analysis/benchmark_results.xlsx")),
        help="Output .xlsx path.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    results_dir = Path(args.results_dir).resolve()
    output_path = Path(args.output).resolve()
    rows = load_all_rows(results_dir)
    _build_base_workbook(rows, output_path)
    print(f"Wrote workbook with {len(rows)} rows to {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
