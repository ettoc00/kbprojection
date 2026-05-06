import argparse
import asyncio
import json
import re
import sys
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.append(str(PROJECT_ROOT))


from kbprojection.llm import AsyncGenericAIClient, DEFAULT_MODELS, LLMGenerationError
from kbprojection.models import NLIProblem, NLILabel
from kbprojection.prompts import fill_prompt
from kbprojection.settings import get_default_results_dir


CURATED_PAIR_PATTERN = re.compile(r"\(\s*([^,()]+?)\s*,\s*([^()]+?)\s*\)")
KB_RELATION_PATTERN = re.compile(r"^\s*([A-Za-z_][A-Za-z0-9_]*)\s*\(\s*([^,()]+?)\s*,\s*([^()]+?)\s*\)\s*$")
PREDICATE_PATTERN = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
LASHA_ANSWER_PATTERN = re.compile(r"^\s*answer:\s*(entailment|non-entailment)\s*$", re.IGNORECASE | re.MULTILINE)
WORD_PATTERN = re.compile(r"[A-Za-z0-9]+(?:'[A-Za-z0-9]+)?")


@dataclass(frozen=True)
class CuratedItem:
    row_number: int
    problem: NLIProblem
    g_value: Any
    curated_kb: list[str]
    raw_kb_cell: str
    data_split: str
    added_by: str


def _normalise_space(value: str) -> str:
    return " ".join(str(value).strip().lower().split())


def validate_predicate(predicate: str) -> str:
    predicate = predicate.strip()
    if not PREDICATE_PATTERN.fullmatch(predicate):
        raise ValueError(f"Invalid predicate name: {predicate}")
    return predicate


def format_relation(predicate: str, left: str, right: str) -> str:
    return f"{predicate}({_normalise_space(left)}, {_normalise_space(right)})"


def normalise_relation(relation: str) -> str:
    match = KB_RELATION_PATTERN.fullmatch(relation.strip())
    if not match:
        raise ValueError(f"Invalid KB relation: {relation}")
    predicate, left, right = match.groups()
    return format_relation(predicate, left, right)


def parse_curated_kb(value: Any, *, predicate: str) -> list[str]:
    if value is None:
        return []
    text = str(value).strip()
    relations = [
        format_relation(predicate, left, right)
        for left, right in CURATED_PAIR_PATTERN.findall(text)
    ]
    if not relations and text:
        raise ValueError(f"Could not parse curated KB cell: {text}")
    return relations


def parse_model_kb(output: str, *, predicate: str) -> list[str]:
    answer_match = LASHA_ANSWER_PATTERN.search(output)
    if not answer_match:
        raise ValueError("Missing 'answer:' line in lasha output.")
    if answer_match.group(1).strip().lower() == "non-entailment":
        return []

    lines = output.splitlines()
    relations_index = None
    for index, line in enumerate(lines):
        if line.strip().lower().startswith("relations:"):
            relations_index = index
            break
    if relations_index is None:
        raise ValueError("Missing 'relations:' line in lasha entailment output.")

    relations_line = lines[relations_index].strip()
    relations_payload = relations_line.split(":", 1)[1].strip()
    if relations_payload.startswith("{") and not relations_payload.endswith("}"):
        payload_lines = [relations_payload]
        for line in lines[relations_index + 1:]:
            payload_lines.append(line.strip())
            if line.strip().endswith("}"):
                break
        relations_payload = " ".join(payload_lines)

    if relations_payload in {"{ }", "{}"}:
        return []
    if not (relations_payload.startswith("{") and relations_payload.endswith("}")):
        raise ValueError(f"Malformed lasha relations block: {relations_line}")

    relations = []
    relation_pattern = re.compile(
        rf"(?:entails|{re.escape(predicate)})\(\s*([^,()]+?)\s*,\s*([^()]+?)\s*\)",
        re.IGNORECASE,
    )
    for left, right in relation_pattern.findall(relations_payload):
        relations.append(format_relation(predicate, left, right))

    if not relations and "(" in relations_payload:
        raise ValueError(f"Could not parse lasha relations: {relations_line}")
    return relations


def build_lasha_prompt(problem: NLIProblem, *, predicate: str) -> str:
    return fill_prompt(
        "lasha",
        problem.premises,
        problem.hypothesis,
        variables={"predicates": {"entailment": predicate}},
    )


def _is_selected_g(value: Any) -> bool:
    if isinstance(value, bool):
        return value is True
    if isinstance(value, (int, float)):
        return float(value) == 1.0
    return str(value).strip() in {"1", "1.0"}


def load_curated_items(path: Path, *, g_only: bool = True, predicate: str = "isa_wn") -> list[CuratedItem]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    headers = [sheet.cell(1, column).value for column in range(1, sheet.max_column + 1)]
    header_index = {name: idx + 1 for idx, name in enumerate(headers) if name}
    required = ["ID", "Premise", "Hypothesis", "G", "KB", "data split", "added by"]
    missing = [name for name in required if name not in header_index]
    if missing:
        raise ValueError(f"Workbook is missing required columns: {missing}")

    items: list[CuratedItem] = []
    for row_number in range(2, sheet.max_row + 1):
        row = {
            name: sheet.cell(row_number, header_index[name]).value
            for name in required
        }
        if not any(value is not None for value in row.values()):
            continue
        if g_only and not _is_selected_g(row["G"]):
            continue

        raw_id = row["ID"]
        if raw_id is None:
            raise ValueError(f"Row {row_number} has no ID")
        if isinstance(raw_id, (int, float)) and int(raw_id) == raw_id:
            problem_id = str(int(raw_id))
        else:
            problem_id = str(raw_id).strip()

        split = str(row["data split"] or "").strip()
        normalised_split = "dev" if split == "trial" else split
        problem = NLIProblem(
            id=problem_id,
            premises=[str(row["Premise"] or "").strip()],
            hypothesis=str(row["Hypothesis"] or "").strip(),
            gold_label=NLILabel.ENTAILMENT,
            dataset="curated_excel",
            split=normalised_split,
            original_data={
                "excel_row": row_number,
                "g": row["G"],
                "data_split": split,
                "added_by": row["added by"],
            },
        )
        items.append(
            CuratedItem(
                row_number=row_number,
                problem=problem,
                g_value=row["G"],
                curated_kb=parse_curated_kb(row["KB"], predicate=predicate),
                raw_kb_cell=str(row["KB"] or ""),
                data_split=split,
                added_by=str(row["added by"] or ""),
            )
        )
    return items


def relation_scores(predicted: list[str], gold: list[str]) -> dict[str, Any]:
    return multiset_scores(predicted, gold)


def multiset_scores(predicted: list[str], gold: list[str]) -> dict[str, Any]:
    predicted_counter = Counter(predicted)
    gold_counter = Counter(gold)
    true_positive = sum((predicted_counter & gold_counter).values())
    predicted_total = sum(predicted_counter.values())
    gold_total = sum(gold_counter.values())
    precision = true_positive / predicted_total if predicted_total else (1.0 if gold_total == 0 else 0.0)
    recall = true_positive / gold_total if gold_total else 1.0
    f1 = (
        2 * precision * recall / (precision + recall)
        if precision + recall
        else 0.0
    )
    return {
        "true_positive": true_positive,
        "predicted_total": predicted_total,
        "gold_total": gold_total,
        "precision": precision,
        "recall": recall,
        "f1": f1,
        "exact_match": predicted_counter == gold_counter,
        "missing": list((gold_counter - predicted_counter).elements()),
        "extra": list((predicted_counter - gold_counter).elements()),
    }


def relation_arguments(relation: str) -> tuple[str, str]:
    match = KB_RELATION_PATTERN.fullmatch(relation.strip())
    if not match:
        raise ValueError(f"Invalid KB relation: {relation}")
    _predicate, left, right = match.groups()
    return left, right


def relation_words(relations: list[str]) -> list[str]:
    words = []
    for relation in relations:
        for argument in relation_arguments(relation):
            words.extend(WORD_PATTERN.findall(_normalise_space(argument)))
    return words


def prefixed_scores(prefix: str, scores: dict[str, Any]) -> dict[str, Any]:
    return {f"{prefix}_{key}": value for key, value in scores.items()}


def score_prediction(predicted: list[str], gold: list[str]) -> dict[str, Any]:
    relation_level = relation_scores(predicted, gold)
    word_level = multiset_scores(relation_words(predicted), relation_words(gold))
    return {
        **relation_level,
        **prefixed_scores("relation", relation_level),
        **prefixed_scores("word", word_level),
    }


async def generate_one(
    client: AsyncGenericAIClient,
    item: CuratedItem,
    *,
    model: str | None,
    predicate: str,
    max_tokens: int | None,
    max_retries: int,
    request_timeout: float | None,
) -> dict[str, Any]:
    prompt = build_lasha_prompt(item.problem, predicate=predicate)
    last_error = ""
    raw_output = ""
    for attempt in range(max_retries + 1):
        try:
            generation = client.generate(prompt=prompt, model=model, max_tokens=max_tokens)
            raw_output = (
                await asyncio.wait_for(generation, timeout=request_timeout)
                if request_timeout
                else await generation
            )
            predicted_kb = parse_model_kb(raw_output, predicate=predicate)
            scores = score_prediction(predicted_kb, item.curated_kb)
            return {
                "row_number": item.row_number,
                "problem": item.problem.model_dump(mode="json"),
                "data_split": item.data_split,
                "added_by": item.added_by,
                "g": item.g_value,
                "curated_kb_cell": item.raw_kb_cell,
                "curated_kb": item.curated_kb,
                "llm_output_raw": raw_output,
                "predicted_kb": predicted_kb,
                "parse_error": "",
                **scores,
            }
        except Exception as exc:
            last_error = str(exc)
            if attempt >= max_retries:
                best_effort = []
                if raw_output:
                    try:
                        best_effort = parse_model_kb(raw_output, predicate=predicate)
                    except Exception:
                        best_effort = []
                scores = score_prediction(best_effort, item.curated_kb)
                return {
                    "row_number": item.row_number,
                    "problem": item.problem.model_dump(mode="json"),
                    "data_split": item.data_split,
                    "added_by": item.added_by,
                    "g": item.g_value,
                    "curated_kb_cell": item.raw_kb_cell,
                    "curated_kb": item.curated_kb,
                    "llm_output_raw": raw_output,
                    "predicted_kb": best_effort,
                    "parse_error": last_error,
                    **scores,
                }
            await asyncio.sleep(min(8.0, 0.5 * (2 ** attempt)))
    raise LLMGenerationError(last_error)


def build_summary(rows: list[dict[str, Any]]) -> dict[str, Any]:
    relation_summary = build_metric_summary(rows, "relation")
    word_summary = build_metric_summary(rows, "word")
    return {
        "rows": len(rows),
        "exact_match_rows": sum(1 for row in rows if row["exact_match"]),
        "parse_error_rows": sum(1 for row in rows if row["parse_error"]),
        "true_positive": relation_summary["true_positive"],
        "predicted_total": relation_summary["predicted_total"],
        "gold_total": relation_summary["gold_total"],
        "micro_precision": relation_summary["micro_precision"],
        "micro_recall": relation_summary["micro_recall"],
        "micro_f1": relation_summary["micro_f1"],
        "macro_precision": relation_summary["macro_precision"],
        "macro_recall": relation_summary["macro_recall"],
        "macro_f1": relation_summary["macro_f1"],
        **prefixed_scores("relation", relation_summary),
        **prefixed_scores("word", word_summary),
    }


def build_metric_summary(rows: list[dict[str, Any]], prefix: str) -> dict[str, Any]:
    tp = sum(row[f"{prefix}_true_positive"] for row in rows)
    predicted_total = sum(row[f"{prefix}_predicted_total"] for row in rows)
    gold_total = sum(row[f"{prefix}_gold_total"] for row in rows)
    precision = tp / predicted_total if predicted_total else 0.0
    recall = tp / gold_total if gold_total else 0.0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
    return {
        "true_positive": tp,
        "predicted_total": predicted_total,
        "gold_total": gold_total,
        "micro_precision": precision,
        "micro_recall": recall,
        "micro_f1": f1,
        "macro_precision": sum(row[f"{prefix}_precision"] for row in rows) / len(rows) if rows else 0.0,
        "macro_recall": sum(row[f"{prefix}_recall"] for row in rows) / len(rows) if rows else 0.0,
        "macro_f1": sum(row[f"{prefix}_f1"] for row in rows) / len(rows) if rows else 0.0,
    }


def load_checkpoint_rows(path: Path) -> dict[int, dict[str, Any]]:
    if not path.exists():
        return {}

    rows: dict[int, dict[str, Any]] = {}
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if not line:
                continue
            row = json.loads(line)
            rows[int(row["row_number"])] = row
    return rows


def append_checkpoint_row(path: Path, row: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(row, ensure_ascii=False) + "\n")


async def run(args: argparse.Namespace) -> int:
    input_path = Path(args.input).resolve()
    output_path = Path(args.output).resolve()
    checkpoint_path = (
        Path(args.checkpoint).resolve()
        if args.checkpoint
        else output_path.with_name(output_path.name + ".checkpoint.jsonl")
    )
    predicate = validate_predicate(args.predicate)
    items = load_curated_items(input_path, g_only=True, predicate=predicate)
    if args.limit is not None:
        items = items[: args.limit]

    checkpoint_rows = load_checkpoint_rows(checkpoint_path) if args.resume else {}
    if args.retry_parse_errors:
        checkpoint_rows = {
            row_number: row
            for row_number, row in checkpoint_rows.items()
            if not row.get("parse_error")
        }
    if args.retry_empty_kb:
        checkpoint_rows = {
            row_number: row
            for row_number, row in checkpoint_rows.items()
            if row.get("predicted_kb")
        }
    if checkpoint_rows:
        print(f"Loaded {len(checkpoint_rows)} checkpointed rows from {checkpoint_path}")

    pending_items = [item for item in items if item.row_number not in checkpoint_rows]

    client = AsyncGenericAIClient(provider=args.provider)
    model = args.model or DEFAULT_MODELS[client.provider]
    semaphore = asyncio.Semaphore(args.concurrency)
    checkpoint_lock = asyncio.Lock()

    async def guarded(item: CuratedItem) -> dict[str, Any]:
        async with semaphore:
            print(f"Generating row {item.row_number} ({item.problem.id})")
            row = await generate_one(
                client,
                item,
                model=model,
                predicate=predicate,
                max_tokens=args.max_tokens,
                max_retries=args.max_retries,
                request_timeout=args.request_timeout,
            )
            async with checkpoint_lock:
                append_checkpoint_row(checkpoint_path, row)
            return row

    generated_rows = await asyncio.gather(*(guarded(item) for item in pending_items))
    rows_by_number = {**checkpoint_rows}
    rows_by_number.update({int(row["row_number"]): row for row in generated_rows})
    rows = [rows_by_number[item.row_number] for item in items if item.row_number in rows_by_number]
    payload = {
        "meta": {
            "input": str(input_path),
            "prompt_style": "lasha",
            "provider": client.provider,
            "model": model,
            "comparison_predicate": predicate,
            "g_filter": 1,
            "langpro_used": False,
            "checkpoint": str(checkpoint_path),
            "checkpoint_rows_loaded": len(checkpoint_rows),
        },
        "summary": build_summary(rows),
        "rows": rows,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(payload["summary"], indent=2))
    print(f"Wrote {len(rows)} evaluated rows to {output_path}")
    return 0


def parse_args() -> argparse.Namespace:
    default_input = Path.home() / "Downloads" / "SICK_KB_test.xlsx"
    default_output = get_default_results_dir() / "curated_excel_lasha_eval.json"
    parser = argparse.ArgumentParser(
        description="Evaluate Lasha-prompt LLM KB generation against curated Excel KB rows with G=1, without LangPro."
    )
    parser.add_argument("--input", default=str(default_input), help="Curated .xlsx file.")
    parser.add_argument("--output", default=str(default_output), help="Output JSON path.")
    parser.add_argument("--provider", default=None, help="Provider override: openai, openrouter, gemini, or claude.")
    parser.add_argument("--model", default=None, help="Model override. Defaults to provider default.")
    parser.add_argument(
        "--predicate",
        default="isa_wn",
        help="Predicate name used in curated/predicted comparison output. Use 'entails' for Lasha-native scoring.",
    )
    parser.add_argument("--limit", type=int, default=None, help="Evaluate only the first N selected rows.")
    parser.add_argument("--concurrency", type=int, default=4, help="Concurrent LLM calls.")
    parser.add_argument("--max-tokens", type=int, default=None, help="Optional provider max token limit.")
    parser.add_argument("--max-retries", type=int, default=2, help="Retries per row after parse/API errors.")
    parser.add_argument(
        "--request-timeout",
        type=float,
        default=180.0,
        help="Seconds before an individual LLM request attempt is treated as failed. Use 0 to disable.",
    )
    parser.add_argument("--checkpoint", default=None, help="Optional JSONL checkpoint path. Defaults next to output.")
    parser.add_argument(
        "--no-resume",
        action="store_false",
        dest="resume",
        help="Ignore any existing checkpoint rows and start this output from scratch.",
    )
    parser.add_argument(
        "--retry-parse-errors",
        action="store_true",
        help="When resuming, rerun checkpointed rows whose previous result has a parse_error.",
    )
    parser.add_argument(
        "--retry-empty-kb",
        action="store_true",
        help="When resuming, rerun checkpointed rows whose previous result has an empty predicted_kb.",
    )
    parser.set_defaults(resume=True)
    return parser.parse_args()


def main() -> int:
    return asyncio.run(run(parse_args()))


if __name__ == "__main__":
    raise SystemExit(main())
